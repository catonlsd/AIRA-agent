# File: backend/tests/test_artifact_enrichment.py
"""Artifact enrichment: richer LLM content (substantive bullets + speaker notes),
bounded web-research grounding when there are no uploaded docs, and safe, optional
image sourcing. Every path stays honest and never breaks generation."""

import json
from pathlib import Path

import pytest

import app.core.llm as llm_module
from app.artifacts.image_providers import OpenverseImageProvider
from app.artifacts.service import ArtifactService
from app.assistant_supervisor import AssistantSupervisor
from app.context_builder import build_turn_context


# ── Richer content: substantive bullets + speaker notes ──────────────────────


def _rich_pptx_json():
    return json.dumps({
        "subtitle": "A grounded overview",
        "slides": [
            {
                "title": "Clay Soil",
                "bullets": [
                    "Clay soil is made of very fine particles that pack tightly and hold water well.",
                    "Its high nutrient retention supports crops but can cause poor drainage and compaction.",
                ],
                "notes": "Clay holds nutrients but drains slowly; amend with organic matter for aeration.",
                "image_query": "clay soil texture close up",
            }
        ],
    })


def test_pptx_content_carries_substantive_bullets_and_notes(monkeypatch):
    monkeypatch.setattr(
        llm_module.LLMClient, "generate",
        lambda self, system, prompt, temperature=0.2: _rich_pptx_json(),
    )
    service = ArtifactService()
    pending, _ = service.plan(
        "Make a PPT on soil types",
        "pptx",
        generate=lambda **kw: llm_module.LLMClient().generate(**kw),
    )
    slides = pending.spec["slides"]
    content = next(s for s in slides if s["title"] == "Clay Soil")
    # Bullets are real sentences, not 1-2 word fragments.
    assert any(len(b.split()) >= 8 for b in content["bullets"])
    assert content["notes"].strip()  # speaker notes preserved


def test_speaker_notes_are_written_into_the_deck(tmp_path):
    from app.artifacts.generators import PptxArtifactGenerator
    from app.artifacts.spec import ArtifactSpec, Slide
    from app.artifacts.styles import get_style

    spec = ArtifactSpec(
        kind="pptx",
        title="Soil",
        slides=[
            Slide(title="Soil", layout="title"),
            Slide(title="Clay", bullets=["Fine particles hold water."], notes="Amend clay with compost.", layout="content"),
        ],
    )
    out = tmp_path / "deck.pptx"
    PptxArtifactGenerator().write(spec, out, get_style("pptx"))
    from pptx import Presentation

    prs = Presentation(str(out))
    notes = [s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide]
    assert any("Amend clay" in n for n in notes)


# ── Web-research grounding when there are no uploaded documents ───────────────


class _ResearchStub:
    def __init__(self):
        self.calls = 0

    def run(self, goal, **kwargs):
        self.calls += 1
        return {
            "message": (
                "Researched facts: loam is roughly 40% sand, 40% silt, and 20% clay, "
                "which gives it balanced drainage and high fertility, making it the "
                "most agriculturally productive soil type for a wide range of crops."
            )
        }


def test_artifact_grounds_in_web_research_when_no_documents(monkeypatch):
    from app.core.config import settings

    monkeypatch.setattr(settings, "artifact_research_grounding", True)

    captured = {}

    def _capture(self, system, prompt, temperature=0.2):
        captured["prompt"] = prompt
        return _rich_pptx_json()

    monkeypatch.setattr(llm_module.LLMClient, "generate", _capture)

    supervisor = AssistantSupervisor()
    supervisor.research = _ResearchStub()
    ctx = build_turn_context("Make a PPT on soil types", session_id="g", owner="g")
    supervisor._present_artifact_plan("Make a PPT on soil types", _Classification(), "pptx", ctx)

    assert supervisor.research.calls == 1  # web research ran to ground content
    # The researched facts were fed into the content planner's prompt.
    assert "Researched facts" in captured["prompt"]


def test_grounding_is_skipped_when_disabled(monkeypatch):
    from app.core.config import settings

    monkeypatch.setattr(settings, "artifact_research_grounding", False)
    monkeypatch.setattr(llm_module.LLMClient, "generate", lambda self, system, prompt, temperature=0.2: _rich_pptx_json())

    supervisor = AssistantSupervisor()
    supervisor.research = _ResearchStub()
    ctx = build_turn_context("Make a PPT on soil types", session_id="g2", owner="g2")
    supervisor._present_artifact_plan("Make a PPT on soil types", _Classification(), "pptx", ctx)
    assert supervisor.research.calls == 0  # no web call when grounding is off


class _Classification:
    mode = "research_then_execution"
    confidence = "high"
    reason = "artifact request"
    artifact_type = "pptx"


# ── Safe image sourcing (fully injectable + guarded) ─────────────────────────


def test_image_provider_returns_path_on_success(tmp_path):
    png = tmp_path / "img.png"
    png.write_bytes(
        bytes.fromhex(
            "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4"
            "890000000a49444154789c63000100000500010d0a2db40000000049454e44ae426082"
        )
    )
    provider = OpenverseImageProvider(
        search=lambda q: "https://example.test/clay.png",
        download=lambda url: str(png),
    )
    assert provider("clay soil") == str(png)


def test_image_provider_caches_per_query():
    calls = {"n": 0}

    def _search(q):
        calls["n"] += 1
        return "https://example.test/x.png"

    provider = OpenverseImageProvider(search=_search, download=lambda url: "/tmp/x.png")
    provider("loam")
    provider("loam")
    assert calls["n"] == 1  # second lookup served from cache


def test_image_provider_is_safe_on_failure():
    # Search raises -> None (never propagates, never breaks generation).
    boom = OpenverseImageProvider(search=lambda q: (_ for _ in ()).throw(RuntimeError("net")))
    assert boom("anything") is None
    # No result / empty query -> None.
    assert OpenverseImageProvider(search=lambda q: None)("x") is None
    assert OpenverseImageProvider(search=lambda q: "u", download=lambda u: None)("x") is None
    assert OpenverseImageProvider()("") is None


def test_build_default_provider_respects_config(monkeypatch):
    from app.artifacts.image_providers import build_default_provider
    from app.core.config import settings

    monkeypatch.setattr(settings, "enable_artifact_images", False)
    assert build_default_provider() is None
    monkeypatch.setattr(settings, "enable_artifact_images", True)
    monkeypatch.setattr(settings, "artifact_image_provider", "openverse")
    assert isinstance(build_default_provider(), OpenverseImageProvider)


def test_deck_with_image_provider_inserts_and_reports(monkeypatch, tmp_path):
    png = tmp_path / "p.png"
    png.write_bytes(
        bytes.fromhex(
            "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4"
            "890000000a49444154789c63000100000500010d0a2db40000000049454e44ae426082"
        )
    )
    import app.artifacts.images as images
    from app.core.config import settings

    monkeypatch.setattr(settings, "enable_artifact_images", True)
    monkeypatch.setattr(settings, "max_artifact_images", 4)
    monkeypatch.setattr(images, "_provider", OpenverseImageProvider(search=lambda q: "u", download=lambda u: str(png)))

    service = ArtifactService()
    pending, _ = service.plan("Make a PPT on soil with pictures", "pptx")
    pending.spec["slides"][-1]["image_query"] = "soil"  # a slide that wants an image
    out = service.generate(pending)
    assert out["status"] == "completed"
    assert out["artifact"]["image_count"] >= 1
    assert out["artifact"]["validation"]["valid"] is True
