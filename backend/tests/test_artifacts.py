# File: backend/tests/test_artifacts.py
"""Artifact pipeline: real PPTX/DOCX/XLSX generation, validation, save-location
approval, and supervisor integration (plan -> approve -> generate -> validate)."""

import pytest

import app.assistant_supervisor as sup_module
from app.artifacts.service import ArtifactService, artifact_store
from app.artifacts.spec import ArtifactPlanBuilder
from app.artifacts.validator import ArtifactValidator
from app.assistant_supervisor import AssistantSupervisor
from app.context_builder import build_turn_context
from app.turn_classifier import RESEARCH_THEN_EXECUTION_MODE, classify_turn

_SESSION = "artifact-session"


@pytest.fixture(autouse=True)
def _artifacts_dir(monkeypatch, tmp_path):
    from app.core.config import settings

    monkeypatch.setattr(settings, "artifacts_dir", str(tmp_path / "artifacts"))
    artifact_store.clear(_SESSION)
    artifact_store.clear(None)
    yield
    artifact_store.clear(_SESSION)
    artifact_store.clear(None)


# ── Routing ──────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "prompt,kind",
    [
        ("Make me a PPT on renewable energy", "pptx"),
        ("Create a DOCX report on AI in healthcare", "docx"),
        ("Generate an XLSX sales tracker", "xlsx"),
        ("Build a presentation from the uploaded files", "pptx"),
    ],
)
def test_artifact_requests_route_with_type(prompt, kind):
    c = classify_turn(prompt)
    assert c.mode == RESEARCH_THEN_EXECUTION_MODE
    assert c.artifact_type == kind


def test_non_artifact_requests_have_no_artifact_type():
    assert classify_turn("what is python?").artifact_type is None
    assert classify_turn("git push").artifact_type is None
    assert classify_turn("summarize this document").artifact_type is None


# ── Real generation + validation ─────────────────────────────────────────────


def _build_and_generate(kind, goal):
    service = ArtifactService()
    pending, _msg = service.plan(goal, kind)  # no LLM -> deterministic fallback
    return service.generate(pending)


def test_pptx_generation_creates_valid_file():
    out = _build_and_generate("pptx", "Make a PPT on renewable energy")
    assert out["status"] == "completed"
    art = out["artifact"]
    assert art["type"] == "pptx"
    from pathlib import Path

    path = Path(art["path"])
    assert path.exists() and path.suffix == ".pptx"
    assert art["validation"]["valid"] is True
    assert art["validation"]["details"]["slides"] >= 1
    # Independently re-open it to prove it's a real presentation.
    from pptx import Presentation

    assert len(Presentation(str(path)).slides) >= 1


def test_docx_generation_creates_valid_file():
    out = _build_and_generate("docx", "Create a report on AI in healthcare")
    assert out["status"] == "completed"
    from pathlib import Path

    from docx import Document

    path = Path(out["artifact"]["path"])
    assert path.exists() and path.suffix == ".docx"
    assert out["artifact"]["validation"]["details"]["paragraphs"] >= 1
    assert len(Document(str(path)).paragraphs) >= 1


def test_xlsx_generation_creates_valid_file():
    out = _build_and_generate("xlsx", "Generate an XLSX sales tracker")
    assert out["status"] == "completed"
    from pathlib import Path

    from openpyxl import load_workbook

    path = Path(out["artifact"]["path"])
    assert path.exists() and path.suffix == ".xlsx"
    wb = load_workbook(str(path))
    assert wb.active.max_row >= 1


def test_metadata_carries_download_url_and_validation():
    out = _build_and_generate("pptx", "Make a PPT on space")
    art = out["artifact"]
    # Download URL is now owner-scoped (default owner "shared" in this helper).
    assert art["download_url"] == f"/artifacts/shared/{art['filename']}"
    assert art["location"] == "workspace"
    assert art["validation"]["type_matches"] is True


# ── Validation honesty ───────────────────────────────────────────────────────


def test_validator_rejects_missing_file(tmp_path):
    res = ArtifactValidator().validate("pptx", tmp_path / "nope.pptx")
    assert res["valid"] is False
    assert "not written" in res["error"]


def test_validator_rejects_wrong_extension(tmp_path):
    bad = tmp_path / "thing.txt"
    bad.write_text("not a pptx")
    res = ArtifactValidator().validate("pptx", bad)
    assert res["valid"] is False
    assert "extension" in res["error"]


# ── Save-location approval ───────────────────────────────────────────────────


def test_default_save_location_needs_no_approval():
    service = ArtifactService()
    pending, _ = service.plan("Make a PPT on energy", "pptx")
    assert pending.delivery["location"] == "workspace"
    assert pending.delivery["requires_approval"] is False


def test_external_save_path_is_flagged_for_approval():
    service = ArtifactService()
    pending, message = service.plan(
        "Make a PPT on energy and save it to C:/Users/me/Desktop", "pptx"
    )
    assert pending.delivery["location"] == "external"
    assert pending.delivery["requires_approval"] is True
    assert "C:/Users/me/Desktop" in message


# ── Supervisor integration: plan -> approve -> validated completion ──────────


@pytest.mark.asyncio
async def test_supervisor_plans_then_generates_on_approval(monkeypatch):
    monkeypatch.setattr(
        sup_module.LLMClient, "generate", lambda self, system, prompt, temperature=0.2: ""
    )
    supervisor = AssistantSupervisor()

    # 1. Artifact request -> plan_ready (NOT completed, no fake artifact).
    ctx = build_turn_context("Make me a PPT on renewable energy", session_id=_SESSION, run_id="a1")
    plan = await supervisor._dispatch("Make me a PPT on renewable energy", ctx)
    assert plan["status"] == "plan_ready"
    assert plan["mode"] == RESEARCH_THEN_EXECUTION_MODE
    assert plan["artifacts"] == []
    assert plan["meta"]["approval_required"] is True
    assert artifact_store.get(_SESSION) is not None

    # 2. Approve -> real file generated, validated, completed with metadata.
    ctx2 = build_turn_context("approve", session_id=_SESSION, run_id="a2")
    done = await supervisor._dispatch("approve", ctx2)
    assert done["status"] == "completed"
    assert done["decision"] == "artifact_generated"
    assert len(done["artifacts"]) == 1
    from pathlib import Path

    assert Path(done["artifacts"][0]["path"]).exists()
    assert done["meta"]["artifact"]["validation"]["valid"] is True
    assert "ready below" in done["message"]  # clickable card carries the download
    assert done["meta"]["artifact"]["download_url"]  # url still in metadata
    assert artifact_store.get(_SESSION) is None  # consumed


@pytest.mark.asyncio
async def test_supervisor_reject_does_not_pretend_delivery():
    supervisor = AssistantSupervisor()
    ctx = build_turn_context("Create a DOCX report on climate", session_id=_SESSION, run_id="r1")
    await supervisor._dispatch("Create a DOCX report on climate", ctx)

    ctx2 = build_turn_context("reject", session_id=_SESSION, run_id="r2")
    res = await supervisor._dispatch("reject", ctx2)
    assert res["decision"] == "artifact_rejected"
    assert res["artifacts"] == []
    assert "didn't generate" in res["message"]
    assert artifact_store.get(_SESSION) is None


@pytest.mark.asyncio
async def test_generation_failure_is_honest(monkeypatch):
    """A generator that produces an empty file must fail, not fake success."""
    supervisor = AssistantSupervisor()
    ctx = build_turn_context("Make a PPT on energy", session_id=_SESSION, run_id="f1")
    await supervisor._dispatch("Make a PPT on energy", ctx)

    # Force the validator to see no slides (corrupt/empty generation).
    monkeypatch.setattr(
        ArtifactValidator,
        "validate",
        lambda self, kind, path: {"valid": False, "error": "presentation has no slides"},
    )
    ctx2 = build_turn_context("approve", session_id=_SESSION, run_id="f2")
    res = await supervisor._dispatch("approve", ctx2)
    assert res["status"] == "failed"
    assert res["decision"] == "artifact_generation_failed"
    assert res["artifacts"] == []
    assert "couldn't produce a valid" in res["message"]


# ── LLM content preparation feeds the generator ──────────────────────────────


def test_plan_builder_uses_llm_content_when_available():
    def _gen(system, prompt, temperature=0.4):
        return '{"slides": [{"title": "Solar", "bullets": ["Cheap", "Clean"]}, {"title": "Wind", "bullets": ["Scalable"]}, {"title": "Hydro", "bullets": ["Reliable"]}]}'

    spec = ArtifactPlanBuilder().build("PPT on renewables", "pptx", generate=_gen)
    titles = [s.title for s in spec.slides]
    assert "Solar" in titles and "Wind" in titles


# ── Step 2: content structuring quality ──────────────────────────────────────


def test_pptx_deck_has_intentional_structure():
    def _gen(system, prompt, temperature=0.4):
        slides = ", ".join(
            f'{{"title": "Topic {i}", "bullets": ["a", "b"]}}' for i in range(1, 5)
        )
        return f'{{"subtitle": "An overview", "slides": [{slides}]}}'

    spec = ArtifactPlanBuilder().build("PPT on energy", "pptx", generate=_gen)
    layouts = [s.layout for s in spec.slides]
    assert layouts[0] == "title"             # opens with a title slide
    assert "agenda" in layouts               # agenda when enough content
    assert layouts[-1] == "summary"          # closes with a summary
    assert spec.subtitle == "An overview"
    # Bullet density is capped (no walls of text).
    assert all(len(s.bullets) <= 6 for s in spec.slides)


def test_docx_has_summary_and_conclusion_and_bullets():
    def _gen(system, prompt, temperature=0.4):
        return (
            '{"sections": [{"heading": "Background", "paragraphs": ["p1"], '
            '"bullets": ["point a", "point b"]}]}'
        )

    spec = ArtifactPlanBuilder().build("Report on AI", "docx", generate=_gen)
    headings = [s.heading.lower() for s in spec.sections]
    assert any("summary" in h for h in headings)       # leading exec summary
    assert any("conclusion" in h for h in headings)    # closing section
    background = next(s for s in spec.sections if s.heading == "Background")
    assert background.bullets == ["point a", "point b"]


def test_xlsx_gets_meaningful_sheet_name():
    spec = ArtifactPlanBuilder().build("Generate an XLSX sales tracker", "xlsx")
    assert spec.sheet_name not in ("Sheet1", "Sheet", "")


# ── Step 2: style/template applied ───────────────────────────────────────────


def test_default_style_profiles_resolve():
    from app.artifacts.styles import get_style

    # Named theme presets are the kind defaults (legacy names alias to them).
    assert get_style("pptx").name == "professional_clean"
    assert get_style("docx").name == "executive_brief"
    assert get_style("xlsx").name == "spreadsheet_clean"
    assert get_style("pptx", "presentation_default").name == "professional_clean"


def test_generated_artifact_records_style_and_summary():
    out = _build_and_generate("xlsx", "Generate an XLSX sales tracker")
    art = out["artifact"]
    assert art["style"] == "spreadsheet_clean"
    assert "rows" in art["summary"] and "columns" in art["summary"]
    assert art["size_bytes"] > 0


def test_xlsx_header_formatting_does_not_break_generation():
    out = _build_and_generate("xlsx", "Generate a budget tracker")
    from pathlib import Path

    from openpyxl import load_workbook

    wb = load_workbook(out["artifact"]["path"])
    ws = wb.active
    # Header row is bold (style applied) and a column width was set.
    assert ws["A1"].font.bold is True
    assert ws.column_dimensions["A"].width is not None


# ── Step 2: optional, safe image support ─────────────────────────────────────


def test_pptx_generation_succeeds_without_images():
    # No image provider is configured by default -> text-only, still valid.
    out = _build_and_generate("pptx", "Make a PPT on the ocean")
    assert out["status"] == "completed"
    assert out["artifact"]["validation"]["valid"] is True


def test_pptx_inserts_image_when_one_resolves(monkeypatch, tmp_path):
    # A tiny real PNG so python-pptx can embed it.
    png = tmp_path / "pic.png"
    png.write_bytes(
        bytes.fromhex(
            "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4"
            "890000000a49444154789c63000100000500010d0a2db40000000049454e44ae426082"
        )
    )
    import app.artifacts.images as images

    monkeypatch.setattr(images, "_provider", lambda q: str(png))

    from app.artifacts.generators import PptxArtifactGenerator
    from app.artifacts.spec import ArtifactSpec, Slide
    from app.artifacts.styles import get_style

    spec = ArtifactSpec(
        kind="pptx",
        title="Imagery",
        slides=[
            Slide(title="Imagery", layout="title"),
            Slide(title="With image", bullets=["a"], layout="content", image_query="ocean"),
        ],
    )
    out_path = tmp_path / "deck.pptx"
    PptxArtifactGenerator().write(spec, out_path, get_style("pptx"))
    from pptx import Presentation

    prs = Presentation(str(out_path))
    # The content slide carries an embedded picture shape.
    pics = [sh for slide in prs.slides for sh in slide.shapes if sh.shape_type == 13]
    assert len(pics) == 1


def test_bad_image_path_does_not_break_generation(monkeypatch, tmp_path):
    import app.artifacts.images as images

    monkeypatch.setattr(images, "_provider", lambda q: "/nonexistent/img.png")
    out = _build_and_generate("pptx", "Make a PPT on mountains")
    assert out["status"] == "completed"  # missing image -> graceful text-only
