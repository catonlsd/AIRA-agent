# File: backend/app/artifacts/validator.py
"""
Artifact validation — proof that a generated file is real and usable.

Completion is never claimed without this passing: the file must exist, match
the requested type, re-open structurally via its library, and contain the
expected content (slides / body / sheet rows). The result feeds evidence,
traces, and the user-facing summary.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any


class ArtifactValidator:
    def validate(self, kind: str, path: Path) -> dict[str, Any]:
        result: dict[str, Any] = {
            "valid": False,
            "exists": False,
            "type_matches": False,
            "opens": False,
            "details": {},
        }
        try:
            if not path.exists() or not path.is_file():
                result["error"] = "file was not written to disk"
                return result
            result["exists"] = True
            result["size_bytes"] = path.stat().st_size

            expected_ext = {"pptx": ".pptx", "docx": ".docx", "xlsx": ".xlsx"}.get(kind)
            if expected_ext and path.suffix.lower() != expected_ext:
                result["error"] = f"extension {path.suffix} does not match requested {kind}"
                return result
            result["type_matches"] = True

            check = {
                "pptx": self._check_pptx,
                "docx": self._check_docx,
                "xlsx": self._check_xlsx,
            }.get(kind)
            if check is None:
                result["error"] = f"no validator for kind {kind}"
                return result

            ok, details = check(path)
            result["opens"] = ok
            result["details"] = details
            result["valid"] = ok
            if not ok:
                result["error"] = details.get("error", "structural validation failed")
        except Exception as error:  # corrupt file etc.
            result["error"] = f"validation error: {error}"
        return result

    @staticmethod
    def _check_pptx(path: Path) -> tuple[bool, dict[str, Any]]:
        from pptx import Presentation

        prs = Presentation(str(path))
        slides = len(prs.slides)
        if slides == 0:
            return False, {"error": "presentation has no slides"}
        return True, {"slides": slides}

    @staticmethod
    def _check_docx(path: Path) -> tuple[bool, dict[str, Any]]:
        from docx import Document

        doc = Document(str(path))
        body = [p for p in doc.paragraphs if (p.text or "").strip()]
        if not body:
            return False, {"error": "document has no body content"}
        return True, {"paragraphs": len(body)}

    @staticmethod
    def _check_xlsx(path: Path) -> tuple[bool, dict[str, Any]]:
        from openpyxl import load_workbook

        wb = load_workbook(str(path), read_only=True)
        sheets = wb.sheetnames
        ws = wb[sheets[0]]
        rows = ws.max_row or 0
        wb.close()
        if not sheets or rows == 0:
            return False, {"error": "workbook has no rows"}
        return True, {"sheets": len(sheets), "rows": rows}
